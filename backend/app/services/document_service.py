"""Document upload, text extraction and lifecycle management.

Supported formats in phase 1: .txt, .md, .pdf, .docx, .xlsx. Extracted text is
stored directly on the Document row and consumed by KnowledgeContextService.
No chunking/embeddings yet (see knowledge_context_service.py).

.xlsx files (e.g. a pricing/budget calculator spreadsheet) are flattened into
a plain-text table per sheet so the LLM can read the figures as part of the
company knowledge context. Only .xlsx (modern Excel) is supported; the legacy
.xls format would require a different library.
"""

import os
import uuid

from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.logging import get_logger
from app.models.document import Document
from app.repositories.document_repository import DocumentRepository

logger = get_logger(__name__)

_SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".xlsx"}


class UnsupportedDocumentTypeError(Exception):
    pass


def _extract_text(path: str, extension: str) -> str:
    if extension in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read()

    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if extension == ".docx":
        import docx

        doc = docx.Document(path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)

    if extension == ".xlsx":
        return _extract_xlsx_text(path)

    raise UnsupportedDocumentTypeError(f"Unsupported extension: {extension}")


def _extract_xlsx_text(path: str) -> str:
    from openpyxl import load_workbook

    # data_only=True reads the last calculated value of each formula cell
    # (e.g. totals in a budget calculator) instead of the formula itself.
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                if all(cell is None for cell in row):
                    continue
                rows.append("\t".join("" if cell is None else str(cell) for cell in row))
            if rows:
                sheets.append(f"### Hoja: {sheet.title}\n" + "\n".join(rows))
        return "\n\n".join(sheets)
    finally:
        workbook.close()


class DocumentService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = DocumentRepository(db)

    def list_documents(self) -> list[Document]:
        return self.repo.list(limit=500)

    def list_active_documents(self) -> list[Document]:
        return self.repo.list_active()

    def get(self, document_id: int) -> Document | None:
        return self.repo.get(document_id)

    def upload(self, *, original_filename: str, content_type: str, file_bytes: bytes) -> Document:
        extension = os.path.splitext(original_filename)[1].lower()
        if extension not in _SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentTypeError(
                f"'{extension}' is not supported. Allowed: {', '.join(sorted(_SUPPORTED_EXTENSIONS))}"
            )

        os.makedirs(settings.documents_storage_path, exist_ok=True)
        stored_filename = f"{uuid.uuid4().hex}{extension}"
        storage_path = os.path.join(settings.documents_storage_path, stored_filename)

        with open(storage_path, "wb") as fh:
            fh.write(file_bytes)

        try:
            extracted_text = _extract_text(storage_path, extension)
        except Exception as exc:
            logger.error("document_text_extraction_failed filename=%s error=%s", original_filename, exc)
            extracted_text = None

        document = Document(
            filename=stored_filename,
            original_filename=original_filename,
            content_type=content_type,
            storage_path=storage_path,
            extracted_text=extracted_text,
            is_active=True,
        )
        self.repo.add(document)
        self.repo.commit()
        logger.info("document_uploaded id=%s filename=%s", document.id, original_filename)
        return document

    def delete(self, document_id: int) -> bool:
        document = self.repo.get(document_id)
        if not document:
            return False
        if os.path.exists(document.storage_path):
            try:
                os.remove(document.storage_path)
            except OSError as exc:
                logger.warning("document_file_delete_failed path=%s error=%s", document.storage_path, exc)
        self.repo.delete(document)
        self.repo.commit()
        logger.info("document_deleted id=%s", document_id)
        return True
